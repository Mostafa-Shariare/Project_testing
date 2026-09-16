"""Analytics aggregation, attendance, and report generation."""
from __future__ import annotations

import csv
import io
import time
from collections import defaultdict
from datetime import datetime, timezone
from typing import Any, Optional

from bson import ObjectId

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

from backend.app.config import ROOT_DIR, get_settings

settings = get_settings()

ALERT_TYPES = {
    "phone": ("PHONE",),
    "no_face": ("NO FACE",),
    "eyes_closed": ("EYES CLOSED",),
}

INSTITUTION_NAME = "Attenova Attention Monitor"
INSTITUTION_SUBTITLE = "Attention Analytics & Session Report"

# Characters that trigger Excel formula injection — prefix with apostrophe to neutralize
_FORBIDDEN_NAME_PREFIXES = ("=", "+", "-", "@", "\t", "\n", "\r")


def _sanitize_export_name(name: str) -> str:
    """Prevent Excel/CSV formula injection by prefixing dangerous values."""
    if name and name.startswith(_FORBIDDEN_NAME_PREFIXES):
        return "'" + name
    return name


def _day_key(ts: float) -> str:
    return datetime.fromtimestamp(ts, tz=timezone.utc).strftime("%Y-%m-%d")


def _week_key(ts: float) -> str:
    dt = datetime.fromtimestamp(ts, tz=timezone.utc)
    return dt.strftime("%Y-W%W")


def _month_key(ts: float) -> str:
    return datetime.fromtimestamp(ts, tz=timezone.utc).strftime("%Y-%m")


def parse_date_range(
    from_ts: Optional[float], to_ts: Optional[float]
) -> tuple[float, float]:
    now = time.time()
    end = to_ts if to_ts is not None else now
    start = from_ts if from_ts is not None else end - 30 * 86400
    if start > end:
        start, end = end, start
    return start, end


def session_duration_sec(doc: dict) -> int:
    start = doc.get("start_time") or 0
    end = doc.get("end_time") or time.time()
    return max(0, int(end - start))


def summarize_session(doc: dict) -> dict:
    logs = doc.get("logs", [])
    avg_attention = round(sum(log["attention"] for log in logs) / len(logs)) if logs else 0
    alerts_count = sum(1 for log in logs if log.get("alert"))
    duration = session_duration_sec(doc)
    return {
        "id": str(doc["_id"]),
        "name": doc["name"],
        "roll_number": doc["roll_number"],
        "class_code": doc["class_code"],
        "teacher_username": doc.get("teacher_username"),
        "start_time": doc["start_time"],
        "end_time": doc["end_time"],
        "status": doc.get("status", "completed"),
        "avg_attention": avg_attention,
        "alerts_count": alerts_count,
        "log_count": len(logs),
        "duration_sec": duration,
        "session_status": doc.get("status", "completed"),
        "join_time": doc.get("join_time", doc.get("start_time")),
        "leave_time": doc.get("leave_time", doc.get("end_time")),
    }


def classify_alert(alert: str) -> Optional[str]:
    upper = (alert or "").strip().upper()
    if not upper:
        return None
    for key, needles in ALERT_TYPES.items():
        if any(n in upper for n in needles):
            return key
    return "other"


def alert_breakdown(logs: list) -> dict:
    counts = {k: 0 for k in list(ALERT_TYPES) + ["other"]}
    for log in logs:
        kind = classify_alert(log.get("alert", ""))
        if kind:
            counts[kind] += 1
    return counts


def fetch_sessions(
    sessions_collection,
    teacher_username: str,
    class_code: Optional[str] = None,
    from_ts: Optional[float] = None,
    to_ts: Optional[float] = None,
    limit: int = 500,
) -> list[dict]:
    query: dict = {"teacher_username": teacher_username}
    if class_code:
        query["class_code"] = class_code.strip().upper()
    start, end = parse_date_range(from_ts, to_ts)
    query["start_time"] = {"$gte": start, "$lte": end}
    cursor = sessions_collection.find(query).sort("start_time", -1).limit(limit)
    return list(cursor)


def rank_sessions(summaries: list[dict]) -> list[dict]:
    ranked = sorted(summaries, key=lambda s: s["avg_attention"], reverse=True)
    total = len(ranked)
    return [{**row, "rank": i + 1, "rank_total": total} for i, row in enumerate(ranked)]


def build_overview(
    sessions: list[dict],
    roster_size: int,
) -> dict:
    summaries = [summarize_session(s) for s in sessions]
    ranked = rank_sessions([dict(s) for s in summaries])

    weekly: dict[str, list[int]] = defaultdict(list)
    monthly: dict[str, list[int]] = defaultdict(list)
    daily_class: dict[str, list[int]] = defaultdict(list)

    for s in summaries:
        ts = s["start_time"]
        weekly[_week_key(ts)].append(s["avg_attention"])
        monthly[_month_key(ts)].append(s["avg_attention"])
        daily_class[_day_key(ts)].append(s["avg_attention"])

    def _avg(vals: list[int]) -> int:
        return round(sum(vals) / len(vals)) if vals else 0

    weekly_trend = [
        {"period": k, "avg_attention": _avg(v), "session_count": len(v)}
        for k, v in sorted(weekly.items())
    ]
    monthly_trend = [
        {"period": k, "avg_attention": _avg(v), "session_count": len(v)}
        for k, v in sorted(monthly.items())
    ]
    class_trend = [
        {"date": k, "avg_attention": _avg(v), "session_count": len(v)}
        for k, v in sorted(daily_class.items())
    ]

    best = ranked[:3] if ranked else []
    worst = list(reversed(ranked[-3:])) if ranked else []

    total_alerts = sum(s["alerts_count"] for s in summaries)
    avg_all = _avg([s["avg_attention"] for s in summaries])

    return {
        "session_count": len(summaries),
        "roster_size": roster_size,
        "avg_attention": avg_all,
        "total_alerts": total_alerts,
        "weekly_trend": weekly_trend,
        "monthly_trend": monthly_trend,
        "class_performance_trend": class_trend,
        "ranked_sessions": ranked,
        "best_sessions": best,
        "worst_sessions": worst,
    }


def compare_sessions(doc_a: dict, doc_b: dict, roster_size: int) -> dict:
    sa = summarize_session(doc_a)
    sb = summarize_session(doc_b)

    def _session_metrics(s: dict) -> dict:
        return {
            "duration_sec": s["duration_sec"],
            "session_status": s.get("session_status", "completed"),
            "participation_minutes": round(s["duration_sec"] / 60, 1),
        }

    return {
        "session_a": {
            **sa,
            "session_metrics": _session_metrics(sa),
            "alert_breakdown": alert_breakdown(doc_a.get("logs", [])),
        },
        "session_b": {
            **sb,
            "session_metrics": _session_metrics(sb),
            "alert_breakdown": alert_breakdown(doc_b.get("logs", [])),
        },
        "delta": {
            "avg_attention": sa["avg_attention"] - sb["avg_attention"],
            "alerts_count": sa["alerts_count"] - sb["alerts_count"],
            "duration_sec": sa["duration_sec"] - sb["duration_sec"],
        },
        "roster_size": roster_size,
    }


def build_student_analytics(
    sessions: list[dict],
    roll_number: str,
    roster_sessions: list[dict],
) -> dict:
    roll = roll_number.strip().upper()
    student_sessions = [s for s in sessions if s.get("roll_number") == roll]
    if not student_sessions:
        return {"roll_number": roll, "session_count": 0}

    summaries = [summarize_session(s) for s in student_sessions]
    all_logs = []
    for s in student_sessions:
        all_logs.extend(s.get("logs", []))

    attentions = [s["avg_attention"] for s in summaries]
    trend = [
        {"session_id": s["id"], "start_time": s["start_time"], "avg_attention": s["avg_attention"]}
        for s in sorted(summaries, key=lambda x: x["start_time"])
    ]

    alerts_total = sum(s["alerts_count"] for s in summaries)
    breakdown: dict[str, int] = {k: 0 for k in list(ALERT_TYPES) + ["other"]}
    for s in student_sessions:
        for k, v in alert_breakdown(s.get("logs", [])).items():
            breakdown[k] += v

    participation_sec = sum(s["duration_sec"] for s in summaries)
    high_focus_count = sum(1 for s in summaries if s["avg_attention"] >= 70)
    high_focus_pct = round(100 * high_focus_count / max(1, len(summaries)))

    # Rank within class by average attention across all sessions
    by_roll: dict[str, list[int]] = defaultdict(list)
    for s in roster_sessions:
        sm = summarize_session(s)
        by_roll[sm["roll_number"]].append(sm["avg_attention"])
    roll_avgs = [(r, round(sum(v) / len(v))) for r, v in by_roll.items() if v]
    roll_avgs.sort(key=lambda x: x[1], reverse=True)
    rank = next((i + 1 for i, (r, _) in enumerate(roll_avgs) if r == roll), None)

    return {
        "roll_number": roll,
        "name": summaries[0]["name"],
        "class_code": summaries[0]["class_code"],
        "session_count": len(summaries),
        "avg_attention": round(sum(attentions) / len(attentions)),
        "highest_attention": max(attentions),
        "lowest_attention": min(attentions),
        "total_alerts": alerts_total,
        "alert_breakdown": breakdown,
        "high_focus_percentage": high_focus_pct,
        "attendance_percentage": high_focus_pct,  # Backwards compatibility alias
        "participation_sec": participation_sec,
        "participation_hours": round(participation_sec / 3600, 2),
        "class_rank": rank,
        "class_rank_total": len(roll_avgs),
        "attention_trend": trend,
        "sessions": summaries,
    }


def build_attention_session_analytics(
    sessions: list[dict],
    roster: Optional[list[dict]] = None,
    from_ts: Optional[float] = None,
    to_ts: Optional[float] = None,
) -> dict:
    summaries = [summarize_session(s) for s in sessions]
    roster_list = roster if roster is not None else [{"roll_number": s["roll_number"]} for s in summaries]
    roster_rolls = {r["roll_number"] for r in roster_list}
    roster_size = len(roster_rolls) or 1

    by_student: dict[str, dict] = {}
    for s in summaries:
        roll = s["roll_number"]
        if roll not in by_student:
            by_student[roll] = {
                "roll_number": roll,
                "name": s["name"],
                "high_focus": 0,
                "moderate_drift": 0,
                "low_focus": 0,
                "sessions": 0,
                "total_duration_sec": 0,
                "attention_scores": [],
            }
        rec = by_student[roll]
        rec["sessions"] += 1
        rec["total_duration_sec"] += s["duration_sec"]
        score = s["avg_attention"]
        rec["attention_scores"].append(score)
        if score >= 70:
            rec["high_focus"] += 1
        elif score >= 45:
            rec["moderate_drift"] += 1
        else:
            rec["low_focus"] += 1

    weekly: dict[str, dict] = defaultdict(lambda: {"high_focus": 0, "moderate_drift": 0, "low_focus": 0})
    monthly: dict[str, dict] = defaultdict(lambda: {"high_focus": 0, "moderate_drift": 0, "low_focus": 0})

    for s in summaries:
        wk = _week_key(s["start_time"])
        mo = _month_key(s["start_time"])
        score = s["avg_attention"]
        if score >= 70:
            weekly[wk]["high_focus"] += 1
            monthly[mo]["high_focus"] += 1
        elif score >= 45:
            weekly[wk]["moderate_drift"] += 1
            monthly[mo]["moderate_drift"] += 1
        else:
            weekly[wk]["low_focus"] += 1
            monthly[mo]["low_focus"] += 1

    leaderboard = sorted(
        [
            {
                "roll_number": v["roll_number"],
                "name": v["name"],
                "high_focus": v["high_focus"],
                "moderate_drift": v["moderate_drift"],
                "low_focus": v["low_focus"],
                "sessions": v["sessions"],
                "total_duration_sec": v["total_duration_sec"],
                "avg_attention": round(sum(v["attention_scores"]) / max(1, len(v["attention_scores"]))),
                "high_focus_rate": round(100 * v["high_focus"] / max(1, v["sessions"]), 1),
                "attendance_rate": round(100 * v["high_focus"] / max(1, v["sessions"]), 1),  # Backwards compat alias
            }
            for v in by_student.values()
        ],
        key=lambda x: x["avg_attention"],
        reverse=True,
    )

    total_sessions = len(summaries)
    high_focus_total = sum(1 for s in summaries if s["avg_attention"] >= 70)
    moderate_drift_total = sum(1 for s in summaries if 45 <= s["avg_attention"] < 70)
    low_focus_total = sum(1 for s in summaries if s["avg_attention"] < 45)

    return {
        "from_time": from_ts,
        "to_time": to_ts,
        "roster_size": roster_size,
        "summary": {
            "total_session_records": total_sessions,
            "high_focus_sessions": high_focus_total,
            "moderate_drift_sessions": moderate_drift_total,
            "low_focus_sessions": low_focus_total,
            "high_focus_rate": round(100 * high_focus_total / max(1, total_sessions), 1),
            "attendance_rate": round(100 * high_focus_total / max(1, total_sessions), 1),  # Backwards compat alias
        },
        "weekly": [{"period": k, **v} for k, v in sorted(weekly.items())],
        "monthly": [{"period": k, **v} for k, v in sorted(monthly.items())],
        "students": leaderboard,
        "records": summaries,
    }


# Backwards compatibility alias
build_attendance_analytics = build_attention_session_analytics


def generate_excel_workbook(
    class_code: str,
    overview: dict,
    student_analytics: list[dict],
    attendance: dict,
    from_ts: float,
    to_ts: float,
) -> bytes:
    wb = Workbook()

    # Session Summary
    ws = wb.active
    ws.title = "Session Summary"
    ws.append([INSTITUTION_NAME, INSTITUTION_SUBTITLE])
    ws.append([f"Class: {class_code}", f"Period: {_fmt_range(from_ts, to_ts)}"])
    ws.append([])
    headers = [
        "Rank",
        "Student",
        "Roll",
        "Started",
        "Duration (min)",
        "Avg Attention",
        "Alerts",
        "Status",
    ]
    ws.append(headers)
    for row in overview.get("ranked_sessions", []):
        ws.append(
            [
                row.get("rank"),
                _sanitize_export_name(row.get("name", "")),
                row["roll_number"],
                _fmt_ts(row["start_time"]),
                round(row["duration_sec"] / 60, 1),
                row["avg_attention"],
                row["alerts_count"],
                row.get("status", "completed"),
            ]
        )

    # Student Performance
    ws2 = wb.create_sheet("Student Performance")
    ws2.append(["Roll", "Name", "Sessions", "Avg Attention %", "Rank", "Alerts", "High Focus %"])
    for st in student_analytics:
        ws2.append(
            [
                st["roll_number"],
                _sanitize_export_name(st.get("name", "")),
                st.get("session_count", st.get("sessions", 0)),
                st.get("avg_attention", 0),
                st.get("class_rank"),
                st.get("total_alerts", 0),
                st.get("high_focus_percentage", st.get("high_focus_rate", st.get("attendance_percentage", 0))),
            ]
        )

    # Alerts
    ws3 = wb.create_sheet("Alerts")
    ws3.append(["Roll", "Name", "Phone", "No Face", "Eyes Closed", "Other"])
    for st in student_analytics:
        b = st.get("alert_breakdown", {})
        ws3.append(
            [
                st["roll_number"],
                _sanitize_export_name(st.get("name", "")),
                b.get("phone", 0),
                b.get("no_face", 0),
                b.get("eyes_closed", 0),
                b.get("other", 0),
            ]
        )

    # Attention Breakdown
    ws4 = wb.create_sheet("Attention Breakdown")
    ws4.append(["Roll", "Name", "High Focus", "Moderate Drift", "Low Focus", "Sessions", "High Focus %"])
    for st in attendance.get("students", []):
        ws4.append(
            [
                st["roll_number"],
                _sanitize_export_name(st.get("name", "")),
                st.get("high_focus", 0),
                st.get("moderate_drift", 0),
                st.get("low_focus", 0),
                st.get("sessions", 0),
                st.get("high_focus_rate", 0),
            ]
        )

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            if cell.value:
                cell.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf_report(
    title: str,
    class_code: str,
    overview: dict,
    student_row: Optional[dict],
    attendance: dict,
    from_ts: float,
    to_ts: float,
) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    w, h = letter
    y = h - inch

    c.setFont("Helvetica-Bold", 16)
    c.drawString(inch, y, INSTITUTION_NAME)
    y -= 20
    c.setFont("Helvetica", 11)
    c.drawString(inch, y, INSTITUTION_SUBTITLE)
    y -= 24
    c.setFont("Helvetica-Bold", 13)
    c.drawString(inch, y, title)
    y -= 18
    c.setFont("Helvetica", 10)
    c.drawString(inch, y, f"Class: {class_code}  |  {_fmt_range(from_ts, to_ts)}")
    y -= 28

    c.setFont("Helvetica-Bold", 11)
    c.drawString(inch, y, "Attention Summary")
    y -= 16
    c.setFont("Helvetica", 10)
    lines = [
        f"Sessions analyzed: {overview.get('session_count', 0)}",
        f"Class average attention: {overview.get('avg_attention', 0)}%",
        f"Total alerts: {overview.get('total_alerts', 0)}",
        f"High focus session rate: {attendance.get('summary', {}).get('high_focus_rate', 0)}%",
    ]
    for line in lines:
        c.drawString(inch, y, line)
        y -= 14

    if student_row:
        y -= 10
        c.setFont("Helvetica-Bold", 11)
        c.drawString(inch, y, f"Student: {_sanitize_export_name(student_row.get('name', ''))} ({student_row.get('roll_number')})")
        y -= 16
        c.setFont("Helvetica", 10)
        for line in [
            f"Avg attention: {student_row.get('avg_attention')}%",
            f"Highest: {student_row.get('highest_attention')}%  Lowest: {student_row.get('lowest_attention')}%",
            f"Total alerts: {student_row.get('total_alerts')}",
            f"Class rank: {student_row.get('class_rank')} / {student_row.get('class_rank_total')}",
            f"High focus rate: {student_row.get('high_focus_percentage', student_row.get('attendance_percentage', 0))}%",
        ]:
            c.drawString(inch, y, line)
            y -= 14

    # Simple bar chart for weekly trend
    trend = overview.get("weekly_trend", [])[-8:]
    if trend and y > 2 * inch:
        y -= 10
        c.setFont("Helvetica-Bold", 11)
        c.drawString(inch, y, "Weekly Attention Trend")
        y -= 12
        chart_x = inch
        chart_y = y - 100
        bar_w = 40
        max_val = max(t["avg_attention"] for t in trend) or 1
        for i, t in enumerate(trend):
            bh = (t["avg_attention"] / max_val) * 80
            x = chart_x + i * (bar_w + 12)
            c.setFillColor(colors.HexColor("#5aa0f0"))
            c.rect(x, chart_y, bar_w, bh, fill=1, stroke=0)
            c.setFillColor(colors.black)
            c.setFont("Helvetica", 7)
            c.drawCentredString(x + bar_w / 2, chart_y - 10, t["period"][-5:])
            c.drawCentredString(x + bar_w / 2, chart_y + bh + 4, str(t["avg_attention"]))

    c.showPage()
    c.save()
    return buf.getvalue()


def attention_session_csv_rows(attendance: dict) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([INSTITUTION_NAME, INSTITUTION_SUBTITLE])
    writer.writerow(["Roll", "Name", "High Focus", "Moderate Drift", "Low Focus", "Sessions", "High Focus Rate %"])
    for st in attendance.get("students", []):
        writer.writerow(
            [
                st["roll_number"],
                _sanitize_export_name(st.get("name", "")),
                st.get("high_focus", 0),
                st.get("moderate_drift", 0),
                st.get("low_focus", 0),
                st.get("sessions", 0),
                st.get("high_focus_rate", 0),
            ]
        )
    return output.getvalue()


# Backwards compatibility alias
attendance_csv_rows = attention_session_csv_rows


def _fmt_ts(ts: Optional[float]) -> str:
    if not ts:
        return ""
    return datetime.fromtimestamp(ts, tz=timezone.utc).strftime("%Y-%m-%d %H:%M")


def _fmt_range(start: float, end: float) -> str:
    return f"{_fmt_ts(start)} — {_fmt_ts(end)}"
